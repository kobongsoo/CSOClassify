# -*- coding: utf-8 -*-
#------------------------------------------------------------------
# 평가셋 생성기 — 업무분류 정밀도·재현율 / 보안등급 C 재현율 (각 500건 규모)
#=> D:\분류함 문서에서 층화 표집해, 사람이 정답을 적어 넣을 라벨링 시트(xlsx)를 만든다.
#   도구가 낸 답은 숨김열에 두고, 사람이 정답을 채우면 '집계' 시트가 정밀도·재현율을
#   표본 기준과 모집단 추정(층별 가중치) 두 가지로 계산한다.
#
#   [단계]
#    1) 문서 목록 — 대상 포맷만 모으고 내용 해시로 중복을 걷어낸다
#    2) 엔진 실행 — Rust 판을 --rule-only 로 한 번 돌려 등급·업무분류·추출 본문을 얻는다
#       (프로덕션 경로 그대로. 파이썬 판과 본문·판정이 같다는 것은 파리티 시험이 지킨다)
#    3) 표집 — 엔진 결과로 층을 나누고, 층마다 포맷 비례로 뽑는다(시드 고정)
#    4) 시트 작성 — 업무분류는 기존 150건(사람 라벨)을 그대로 살려 넣고 나머지를 채운다
#
#   [왜 엔진 결과로 층을 나누나] 사람이 정답을 달기 전에 쓸 수 있는 정보가 그것뿐이다.
#   특히 보안등급은 '엔진이 C 가 아니라고 한 문서'를 일부러 많이 넣어야 놓친 C 를
#   찾을 수 있다 — C 로 판정된 문서만 보면 재현율은 절대 잴 수 없다.
#
#   사용: python tests/evalset/build_evalsets.py [--work 작업폴더]
#   같은 입력·같은 시드면 같은 표본이 나온다.
#------------------------------------------------------------------

import argparse
import collections
import hashlib
import json
import math
import os
import random
import re
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC_DIR = r"D:\분류함"
# 1차 개발 대상 10종 + 기존 검증셋에 들어 있던 md·html.
EXT = {".txt", ".hwp", ".hwpx", ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx",
       ".pdf", ".md", ".html"}
SEED = 20260914
MIN_TEXT_LEN = 20
RS_EXE = os.path.join(ROOT, "Rust", "target", "release", "MpowerClassify-rs.exe")
POLICY_DIR = os.path.join(ROOT, "resources", "policy")
OLD_XLSX = os.path.join(ROOT, "report", "업무분류-정밀도검증셋-라벨링시트.xlsx")
OUT_DOCTYPE = os.path.join(ROOT, "report", "업무분류-평가셋-500-20260914.xlsx")
OUT_SECURITY = os.path.join(ROOT, "report", "보안등급-C재현율-평가셋-500-20260914.xlsx")

HDR_FILL = PatternFill("solid", fgColor="FF2F5597")
IN_FILL = PatternFill("solid", fgColor="FFFFF2CC")
HDR_FONT = Font(bold=True, color="FFFFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SEC_FONT = Font(bold=True, size=11, color="FF2F5597")

DOC_STRATA = ("D1_라벨1개", "D2_라벨2개이상", "D3_미라벨")
SEC_STRATA = ("G1_엔진C", "G2_엔진S", "G3_엔진O", "G4_엔진미분류")
SEC_BASIS = "개인정보,법상민감정보,영업·기술기밀,계약·재무,인사·급여,대외비표시,사내한정업무,공개자료,기타"


#------------------------------------------------------------------
# 대상 문서 목록 만들기 (내용 해시로 중복 제거)
#=> 같은 파일이 여러 폴더에 복사돼 있으면 한 문서로 본다. 같은 문서가 표본에 두 번
#   들어가면 사람이 같은 판정을 두 번 하고, 집계에서 그 문서만 두 배로 세어진다.
#    1) 대상 포맷(EXT)만 모아 파일마다 SHA-256 을 구한다
#    2) 해시가 같으면 하나로 묶고, '99_중복파일' 이 아닌 짧은 경로를 대표로 삼는다
#       (폴더 이름이 업무 맥락을 담기 때문)
#
# -in: src_dir = 문서 폴더(기본 D:\분류함)
#
# -out: list[dict] = {path, ext, size, hash, dups, all_paths}
# -out: error = 읽을 수 없는 파일은 OSError 전파
#------------------------------------------------------------------
def enumerate_docs(src_dir):
    rows = []
    for root, _, files in os.walk(src_dir):
        for name in files:
            ext = os.path.splitext(name)[1].lower()
            if ext not in EXT:
                continue
            path = os.path.join(root, name)
            with open(path, "rb") as fp:
                digest = hashlib.sha256(fp.read()).hexdigest()
            rows.append({"path": path, "ext": ext, "size": os.path.getsize(path), "hash": digest})
    by_hash = collections.defaultdict(list)
    for r in rows:
        by_hash[r["hash"]].append(r)
    uniq = []
    for digest in sorted(by_hash):
        group = sorted(by_hash[digest], key=lambda r: ("99_중복파일" in r["path"], len(r["path"]), r["path"]))
        rep = dict(group[0])
        rep["dups"] = len(group)
        rep["all_paths"] = [r["path"] for r in group]
        uniq.append(rep)
    return uniq


#------------------------------------------------------------------
# 엔진 한 번 돌리기 (이미 돌렸으면 건너뜀)
#=> Rust 판을 --filelist 로 모든 대상 문서에 돌려 결과(run.jsonl)와 추출 본문(text/)을
#   남긴다. 수십 분 걸리는 단계라 결과가 완성돼 있으면 다시 돌리지 않는다.
#
# -in: work = 작업폴더
# -in: docs = enumerate_docs 결과
#
# -out: 없음(work 에 filelist.jsonl · run.jsonl · text/ 를 만든다)
# -out: error = exe 가 없으면 SystemExit
#------------------------------------------------------------------
def run_engine(work, docs):
    out = os.path.join(work, "run.jsonl")
    if os.path.isfile(out) and '"summary"' in open(out, encoding="utf-8").read()[-4000:]:
        print(f"[엔진] 기존 결과 사용: {out}")
        return
    if not os.path.isfile(RS_EXE):
        sys.exit(f"Rust exe 가 없습니다: {RS_EXE} (cargo build --release 먼저)")
    flist = os.path.join(work, "filelist.jsonl")
    with open(flist, "w", encoding="utf-8") as fp:
        for d in docs:
            # sfile_id 는 목록 형식이 요구해서 넣는 자리표시값이다(평가에는 쓰지 않는다).
            fp.write(json.dumps({"path": d["path"], "sfile_id": "EV" + d["hash"][:16]},
                                ensure_ascii=False) + "\n")
    env = dict(os.environ, CSOCLASSIFY_POLICY_DIR=POLICY_DIR)
    print("[엔진] 실행 중 — 문서 수에 따라 수십 분 걸릴 수 있습니다")
    with open(os.path.join(work, "run.err"), "w", encoding="utf-8") as err:
        subprocess.run([RS_EXE, "--filelist", flist, "--rule-only", "--textsave", "text",
                        "--format", "jsonl", "--out", "run.jsonl"],
                       cwd=work, env=env, stderr=err, check=False)


#------------------------------------------------------------------
# 경로 정규화 (비교용)
#=> 엔진 결과의 file 과 목록의 path 가 구분자·대소문자만 달라도 같은 문서로 잇는다.
#
# -in: p = 파일 경로
#
# -out: str = 비교용으로 정규화한 경로
# -out: error = 없음
#------------------------------------------------------------------
def norm_path(p):
    return os.path.normcase(os.path.normpath(p))


#------------------------------------------------------------------
# 엔진 결과 읽어 문서 정보에 붙이기
#=> 문서마다 엔진 등급·업무분류·추출 본문을 모은다.
#    1) run.jsonl 의 레코드를 경로로 문서 목록과 잇는다
#    2) 저장된 본문(text/<hash>.txt)을 읽어 공백 제외 글자 수를 센다
#    3) 추출 실패·처리제외·본문 부족 문서는 usable=False — 사람이 발췌로 판정할 수 없다
#
# -in: work = 작업폴더
# -in: docs = enumerate_docs 결과(제자리에서 칸을 더한다)
#
# -out: 없음(docs 의 각 dict 에 rec·text·chars·usable 을 단다)
# -out: error = run.jsonl 이 없으면 FileNotFoundError
#------------------------------------------------------------------
def attach_results(work, docs):
    by_path = {norm_path(d["path"]): d for d in docs}
    for line in open(os.path.join(work, "run.jsonl"), encoding="utf-8"):
        line = line.strip()
        if not line.startswith("{"):
            continue
        rec = json.loads(line)
        if "file" not in rec:
            continue
        d = by_path.get(norm_path(rec["file"]))
        if d is not None:
            d["rec"] = rec
    for d in docs:
        rec = d.get("rec") or {}
        text = ""
        if rec.get("text_saved"):
            tp = os.path.join(work, "text", rec["text_saved"])
            if os.path.isfile(tp):
                text = open(tp, encoding="utf-8").read()
        d["text"] = text
        d["chars"] = len(re.sub(r"\s", "", text))
        d["usable"] = bool(rec) and "error" not in rec and "excluded" not in rec \
            and d["chars"] >= MIN_TEXT_LEN


#------------------------------------------------------------------
# 최대잉여 배분 (정수 몫 나누기)
#=> total 을 가중치 비례로 정수로 나눈다. 반올림만 하면 합이 total 과 어긋나므로
#   내림한 뒤 소수부가 큰 순서로 1 씩 더 준다(같으면 이름순 — 결과가 흔들리지 않게).
#
# -in: total   = 나눌 정수
# -in: weights = {키: 가중치}
#
# -out: dict = {키: 정수 몫} (합 = total, 가중치 합이 0 이면 전부 0)
# -out: error = 없음
#------------------------------------------------------------------
def largest_remainder(total, weights):
    wsum = sum(weights.values())
    if total <= 0 or wsum <= 0:
        return {k: 0 for k in weights}
    quota = {k: total * w / wsum for k, w in weights.items()}
    share = {k: int(math.floor(q)) for k, q in quota.items()}
    left = total - sum(share.values())
    for k in sorted(weights, key=lambda k: (-(quota[k] - share[k]), k))[:left]:
        share[k] += 1
    return share


#------------------------------------------------------------------
# 층별 표본 수 정하기
#=> 층마다 먼저 바닥(floor)과 이미 확보한 표본(fixed) 중 큰 값을 주고, 남은 몫을
#   √(모집단) 비례로 나눈다. 모집단 비례로만 나누면 작은 층이 몇 건밖에 안 뽑혀
#   그 층의 비율을 잴 수 없고, 균등하게 나누면 큰 층이 과소대표된다 — √N 은 그 중간이다.
#   어떤 층이 모집단을 다 써서 더 못 받으면 넘친 몫을 나머지 층에 다시 나눈다.
#
# -in: pop   = {층: 모집단 문서 수}
# -in: total = 전체 표본 목표
# -in: fixed = {층: 이미 확보한 표본 수}(기본 없음)
# -in: floor = 층마다 최소 표본(기본 0, 모집단보다 크면 모집단까지)
#
# -out: dict = {층: 목표 표본 수(fixed 포함)}
# -out: error = 없음
#------------------------------------------------------------------
def allocate(pop, total, fixed=None, floor=0):
    fixed = fixed or {}
    target = {s: min(n, max(floor, fixed.get(s, 0))) for s, n in pop.items()}
    remain = total - sum(target.values())
    while remain > 0:
        open_ = [s for s in pop if target[s] < pop[s]]
        if not open_:
            break
        share = largest_remainder(remain, {s: math.sqrt(pop[s]) for s in open_})
        moved = 0
        for s in open_:
            add = min(share[s], pop[s] - target[s])
            target[s] += add
            moved += add
        remain -= moved
        if moved == 0:
            break
    return target


#------------------------------------------------------------------
# 한 층에서 포맷 비례로 뽑기
#=> 층 안에서도 포맷이 한쪽으로 쏠리지 않게, 포맷별 문서 수에 비례해 몫을 나눈 뒤
#   포맷마다 무작위로 뽑는다. 문서를 해시순으로 정렬한 뒤 뽑아, 파일시스템 나열 순서가
#   달라도 같은 시드면 같은 표본이 나온다.
#
# -in: rng  = random.Random (층마다 고정 시드)
# -in: docs = 그 층의 뽑을 수 있는 문서들
# -in: k    = 뽑을 개수
#
# -out: list[dict] = 뽑힌 문서
# -out: error = 없음(k 가 문서 수보다 크면 전부)
#------------------------------------------------------------------
def sample_by_format(rng, docs, k):
    if k >= len(docs):
        return list(docs)
    groups = collections.defaultdict(list)
    for d in sorted(docs, key=lambda d: d["hash"]):
        groups[d["ext"]].append(d)
    share = largest_remainder(k, {e: len(g) for e, g in groups.items()})
    picked = []
    for ext in sorted(groups):
        picked.extend(rng.sample(groups[ext], min(share[ext], len(groups[ext]))))
    return picked


#------------------------------------------------------------------
# 발췌 만들기
#=> 사람이 판정에 쓸 앞부분 글자. 빈 줄은 빼고 줄 사이를 ' ⏎ ' 로 이어 한 칸에 담는다
#   (기존 검증셋과 같은 표기).
#
# -in: text  = 추출 본문
# -in: limit = 최대 글자 수
#
# -out: (첫 줄, 발췌) = (60자 이내 첫 줄, limit 자 이내 발췌)
# -out: error = 없음
#------------------------------------------------------------------
def excerpt(text, limit):
    lines = [l.strip() for l in (text or "").splitlines() if l.strip()]
    first = lines[0][:60] if lines else ""
    return first, " ⏎ ".join(lines)[:limit]


#------------------------------------------------------------------
# 셀에 글자 넣기 (수식·깨진 문자 방지)
#=> 본문 발췌가 '=' 로 시작하면 엑셀이 수식으로 읽어 오류가 나고, 제어문자가 섞이면
#   파일 저장이 실패한다. 제어문자를 걷어내고 '문자열'로 못박아 넣는다.
#
# -in: ws    = 워크시트
# -in: row   = 행 번호
# -in: col   = 열 번호
# -in: value = 넣을 값(None 이면 빈 칸)
#
# -out: cell = 값이 들어간 셀
# -out: error = 없음
#------------------------------------------------------------------
def put_text(ws, row, col, value):
    cell = ws.cell(row, col)
    if value is None or value == "":
        return cell
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    cell.data_type = "s"
    return cell


#------------------------------------------------------------------
# 표 머리줄 꾸미기
#=> 기존 검증셋과 같은 모양(진한 파랑 바탕·흰 굵은 글자)으로 머리줄을 쓴다.
#
# -in: ws      = 워크시트
# -in: headers = 머리줄 글자 목록
# -in: widths  = {열 글자: 너비}
#
# -out: 없음
# -out: error = 없음
#------------------------------------------------------------------
def write_header(ws, headers, widths):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(1, i, h)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center")
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


#------------------------------------------------------------------
# 안내·집계 시트에 줄 쓰기 (도우미)
#=> 제목(■)은 강조하고, 나머지는 칸마다 값을 넣는다. '=' 로 시작하는 값은 수식으로 둔다.
#
# -in: ws   = 워크시트
# -in: rows = [[값, ...], ...] (빈 목록이면 빈 줄)
#
# -out: 없음
# -out: error = 없음
#------------------------------------------------------------------
def write_rows(ws, rows):
    for r, vals in enumerate(rows, start=1):
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            if c == 1 and isinstance(v, str) and v.startswith("■"):
                cell.font = SEC_FONT
        if r == 1:
            ws.cell(1, 1).font = TITLE_FONT
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16


#------------------------------------------------------------------
# 기존 검증셋 150건을 지금 문서와 잇기
#=> 1차 검증셋(2026-08-26)의 사람 라벨을 버리지 않고 새 평가셋에 그대로 옮긴다.
#   D:\분류함 폴더 구성이 그 뒤로 바뀌어('1.제품매뉴얼' → '01_제품매뉴얼') 경로로는
#   못 잇는다 — 파일 이름으로 찾고, 같은 이름이 여러 내용이면 '첫 줄'로 가린다.
#    1) 이름이 같은 파일(긴 이름은 앞부분 일치)의 내용 해시 후보를 모은다
#    2) 후보가 하나면 그 문서, 여럿이면 첫 줄이 같은 것, 그래도 여럿이면 잇지 않는다
#
# -in: docs = attach_results 까지 끝난 문서 목록
#
# -out: (matched, dropped) = ([{doc, old}], [(파일명, 사유)])
# -out: error = 기존 파일이 없으면 FileNotFoundError
#------------------------------------------------------------------
def match_old_rows(docs):
    wb = openpyxl.load_workbook(OLD_XLSX)
    ws = wb["검증셋"]
    by_name = collections.defaultdict(dict)
    for d in docs:
        for p in d["all_paths"]:
            by_name[os.path.basename(p)][d["hash"]] = d
    matched, dropped = [], []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        old = {"name": ws.cell(r, 3).value, "first": ws.cell(r, 6).value,
               "ans1": ws.cell(r, 8).value, "ans2": ws.cell(r, 9).value,
               "verdict": ws.cell(r, 10).value, "note": ws.cell(r, 11).value,
               "snap": ws.cell(r, 13).value}
        name = str(old["name"] or "")
        cands = dict(by_name.get(name, {}))
        if not cands:
            # 1차 시트는 아주 긴 파일명을 잘라 적었다 — 앞부분이 같으면 같은 파일로 본다.
            for n, hs in by_name.items():
                if name.strip() and n.startswith(name.strip()):
                    cands.update(hs)
        if len(cands) > 1 and old["first"]:
            key = str(old["first"]).strip()[:20]
            cands = {h: d for h, d in cands.items() if excerpt(d["text"], 60)[0][:20] == key}
        if len(cands) != 1:
            dropped.append((name, "파일을 찾지 못함" if not cands else "같은 이름 문서가 여럿"))
            continue
        d = next(iter(cands.values()))
        if not d["usable"]:
            dropped.append((name, "지금 엔진에서 본문을 못 읽음(실패·처리제외)"))
            continue
        matched.append({"doc": d, "old": old})
    # 같은 문서가 1차 시트에 두 번 있었다면 한 번만 쓴다(첫 행의 라벨을 살린다).
    # 1차 모집단은 복사본을 걸러내지 않아 같은 문서가 여러 행으로 뽑혔다. 두 행의 라벨이
    # 다르면 한쪽을 조용히 고르지 않고 비고에 [확인필요] 로 남긴다 — 사람 판정이 갈린 문서다.
    seen, uniq = {}, []
    for m in matched:
        key = m["doc"]["hash"]
        if key in seen:
            kept = seen[key]["old"]
            mine = (m["old"]["ans1"], m["old"]["ans2"], m["old"]["verdict"])
            if mine != (kept["ans1"], kept["ans2"], kept["verdict"]):
                kept["note"] = (f"[확인필요] 1차 시트에 같은 문서가 두 번 있었고 라벨이 달랐다"
                                f"(다른 행: 정답1={mine[0] or '없음'}, 정답2={mine[1] or '없음'}, "
                                f"판정={mine[2] or '없음'}). 둘 중 맞는 쪽으로 고쳐 주세요. "
                                + (kept["note"] or "")).strip()
                dropped.append((m["old"]["name"], "1차 시트 안에서 같은 문서 중복 — 라벨이 달라 [확인필요] 표시"))
            else:
                dropped.append((m["old"]["name"], "1차 시트 안에서 같은 문서 중복(라벨 같음)"))
            continue
        seen[key] = m
        uniq.append(m)
    return uniq, dropped


#------------------------------------------------------------------
# 업무분류 층 가르기
#=> 엔진이 붙인 업무분류 라벨 수로 층을 나눈다. 1차에서 오탐이 '라벨 2개 이상'에 몰렸고,
#   놓침은 '미라벨'에 숨어 있으므로 셋을 따로 재야 한다.
#
# -in: rec = 엔진 결과 레코드
#
# -out: str = DOC_STRATA 중 하나
# -out: error = 없음
#------------------------------------------------------------------
def doctype_stratum(rec):
    n = len(rec.get("doctype") or [])
    return DOC_STRATA[0] if n == 1 else DOC_STRATA[1] if n >= 2 else DOC_STRATA[2]


#------------------------------------------------------------------
# 보안등급 층 가르기
#=> 엔진 등급(C/S/O/미분류)으로 층을 나눈다. 놓친 C 는 C 가 아닌 층에 숨어 있다.
#
# -in: rec = 엔진 결과 레코드
#
# -out: str = SEC_STRATA 중 하나
# -out: error = 없음
#------------------------------------------------------------------
def security_stratum(rec):
    return {"C": SEC_STRATA[0], "S": SEC_STRATA[1], "O": SEC_STRATA[2]}.get(rec.get("grade"), SEC_STRATA[3])


#------------------------------------------------------------------
# 업무분류 근거 신호원 요약
#=> 붙은 라벨들이 본문(body)·표제부(head)·파일명(name) 중 어디서 왔는지 모아 적는다.
#   사람 판정이 끝난 뒤 '어느 신호가 오탐을 냈나'를 가르는 데 쓴다.
#
# -in: rec = 엔진 결과 레코드
#
# -out: str | None = "body;head;name" 모양, 라벨이 없으면 None
# -out: error = 없음
#------------------------------------------------------------------
def doctype_sources(rec):
    vals = ((rec.get("why") or {}).get("doctype") or {}).get("values") or []
    src = sorted({s for v in vals for s in (v.get("from") or [])})
    return ";".join(src) or None


#------------------------------------------------------------------
# 보안등급 근거 신호 요약
#=> 어떤 신호(규칙·민감정보·스탬프·파일명)가 등급을 정했고, 규칙 중 무엇이 몇 번
#   걸렸는지 짧게 적는다. 사람 판정이 끝난 뒤 놓침·오탐의 원인을 가르는 데 쓴다.
#
# -in: rec = 엔진 결과 레코드
#
# -out: str = "rule | IP주소×8, 프로젝트·제품명×3" 모양(최대 6개)
# -out: error = 없음
#------------------------------------------------------------------
def security_signals(rec):
    sec = (rec.get("why") or {}).get("security") or {}
    hits = []
    for sig in (sec.get("signals") or {}).values():
        for h in (sig or {}).get("hits") or []:
            hits.append(f"{h.get('name')}×{h.get('count')}")
    decided = ";".join(sec.get("decided_by") or [])
    return f"{decided} | {', '.join(hits[:6])}" if hits else decided


#------------------------------------------------------------------
# 폴더 표기 (문서 폴더 기준 상대)
#=> 사람이 '어느 폴더에 있던 문서인가'로 맥락을 짐작할 수 있게 상대 폴더를 적는다.
#
# -in: path = 파일 경로
#
# -out: str = D:\분류함 기준 상대 폴더
# -out: error = 없음
#------------------------------------------------------------------
def rel_folder(path):
    return os.path.dirname(os.path.relpath(path, SRC_DIR))


#------------------------------------------------------------------
# 업무분류 평가셋 표집
#=> 1차 150건(사람 라벨)을 먼저 확보하고, 층별 목표에서 모자란 만큼만 새로 뽑는다.
#
# -in: docs  = 문서 목록(usable 판정 끝)
# -in: total = 목표 건수(기본 500)
#
# -out: (rows, pop, dropped) = (표본 행 목록, {층: 모집단 수}, 1차에서 못 옮긴 목록)
# -out: error = 없음
#------------------------------------------------------------------
def pick_doctype(docs, total):
    usable = [d for d in docs if d["usable"]]
    pop_docs = collections.defaultdict(list)
    for d in usable:
        pop_docs[doctype_stratum(d["rec"])].append(d)
    pop = {s: len(pop_docs[s]) for s in DOC_STRATA}
    old, dropped = match_old_rows(docs)
    old_hash = {m["doc"]["hash"] for m in old}
    fixed = collections.Counter(doctype_stratum(m["doc"]["rec"]) for m in old)
    target = allocate(pop, total, fixed=fixed)
    rows = [{"doc": m["doc"], "old": m["old"], "round": "1차"} for m in old]
    new = []
    for s in DOC_STRATA:
        need = target[s] - fixed.get(s, 0)
        cand = [d for d in pop_docs[s] if d["hash"] not in old_hash]
        rng = random.Random(f"{SEED}:doctype:{s}")
        new.extend({"doc": d, "old": None, "round": "2차"} for d in sample_by_format(rng, cand, max(0, need)))
    # 새로 뽑은 문서는 층·포맷이 몰려 나오지 않게 섞는다 — 줄지어 같은 종류가 나오면 판정이 관성에 끌린다.
    random.Random(f"{SEED}:doctype:order").shuffle(new)
    return rows + new, pop, dropped


#------------------------------------------------------------------
# 보안등급 평가셋 표집
#=> 엔진 등급 층마다 바닥 60건을 먼저 주고 나머지를 √N 비례로 나눠 뽑는다.
#   C 층은 정밀도(C 로 판정한 것이 맞나)를, 나머지 층은 재현율(놓친 C)을 잰다.
#
# -in: docs  = 문서 목록(usable 판정 끝)
# -in: total = 목표 건수(기본 500)
#
# -out: (rows, pop) = (표본 행 목록, {층: 모집단 수})
# -out: error = 없음
#------------------------------------------------------------------
def pick_security(docs, total):
    pop_docs = collections.defaultdict(list)
    for d in docs:
        if d["usable"]:
            pop_docs[security_stratum(d["rec"])].append(d)
    pop = {s: len(pop_docs[s]) for s in SEC_STRATA}
    target = allocate(pop, total, floor=60)
    rows = []
    for s in SEC_STRATA:
        rng = random.Random(f"{SEED}:security:{s}")
        rows.extend({"doc": d} for d in sample_by_format(rng, pop_docs[s], target[s]))
    random.Random(f"{SEED}:security:order").shuffle(rows)
    return rows, pop


#------------------------------------------------------------------
# 업무분류 채점 수식 (행 하나)
#=> 1차 검증셋과 같은 네 갈래 판정식이다(설계서 3-C-2). 사람 판정이 비어 있으면 빈칸.
#    · 분류대상아님·체계에없음 → 엔진이 붙였으면 헛라벨, 안 붙였으면 정상무라벨
#    · 정답있음 → 안 붙였으면 놓침, 정답1·정답2 중 하나라도 겹치면 적중, 아니면 오라벨
#
# -in: r   = 행 번호
# -in: col = 엔진 라벨이 든 열 글자("M" 등, 여러 라벨은 ';' 로 이어져 있다)
#
# -out: str = 엑셀 수식
# -out: error = 없음
#------------------------------------------------------------------
def doctype_formula(r, col):
    return (f'=IF($J{r}="","",IF(OR($J{r}="분류대상아님",$J{r}="체계에없음"),'
            f'IF(${col}{r}="","정상무라벨","헛라벨"),IF($H{r}="","",IF(${col}{r}="","놓침",'
            f'IF(OR(ISNUMBER(SEARCH(";"&$H{r}&";",";"&${col}{r}&";")),AND($I{r}<>"",'
            f'ISNUMBER(SEARCH(";"&$I{r}&";",";"&${col}{r}&";")))),"적중","오라벨")))))')


#------------------------------------------------------------------
# 업무분류 평가셋 xlsx 쓰기
#=> 시트 넷: 작성지침 · 평가셋 · 분류체계(드롭다운 원본) · 집계(전부 수식).
#   엔진 답(ref_*)과 채점 결과는 숨김열에 둔다 — 보고 나서 정답을 달면 채점이 아니라
#   베끼기가 된다.
#
# -in: rows    = pick_doctype 의 표본 행
# -in: pop     = {층: 모집단 수}
# -in: dropped = 1차에서 옮기지 못한 목록
# -in: meta    = {rule_version, doctype_rule_version, n_population, ...}
# -in: path    = 저장 경로
#
# -out: 없음(파일을 쓴다)
# -out: error = 저장 실패 시 OSError 전파
#------------------------------------------------------------------
def write_doctype_book(rows, pop, dropped, meta, path):
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "작성지침"
    ws = wb.create_sheet("평가셋")
    tax = wb.create_sheet("분류체계")
    agg = wb.create_sheet("집계")
    last = len(rows) + 1
    n_old = sum(1 for x in rows if x["round"] == "1차")

    headers = ["no", "폴더", "파일명", "포맷", "규모", "첫 줄", "앞부분 발췌", "정답1_dc_id",
               "정답2_dc_id", "판정", "비고", "ref_상대경로", "ref_현행엔진", "ref_1차스냅샷(0826)",
               "ref_신호원", "ref_층", "ref_회차", "결과_현행엔진", "결과_1차스냅샷"]
    write_header(ws, headers, {"A": 5, "B": 20, "C": 42, "D": 7, "E": 10, "F": 34, "G": 60,
                               "H": 17, "I": 14, "J": 14, "K": 26, "L": 30, "M": 26, "N": 26,
                               "O": 14, "P": 16, "Q": 8, "R": 12, "S": 12})
    for i, x in enumerate(rows, start=1):
        r = i + 1
        d, rec, old = x["doc"], x["doc"]["rec"], x["old"] or {}
        first, exc = excerpt(d["text"], 300)
        ws.cell(r, 1, i)
        put_text(ws, r, 2, rel_folder(d["path"]))
        put_text(ws, r, 3, os.path.basename(d["path"]))
        put_text(ws, r, 4, d["ext"])
        put_text(ws, r, 5, f"{d['chars']:,}자")
        put_text(ws, r, 6, first)
        put_text(ws, r, 7, exc).alignment = Alignment(wrap_text=True, vertical="top")
        # 1차 150건은 사람이 이미 단 정답을 그대로 옮긴다(판정·비고까지).
        for col, key in ((8, "ans1"), (9, "ans2"), (10, "verdict"), (11, "note")):
            put_text(ws, r, col, old.get(key))
        for col in range(8, 12):
            ws.cell(r, col).fill = IN_FILL
        put_text(ws, r, 12, os.path.relpath(d["path"], SRC_DIR))
        put_text(ws, r, 13, ";".join(rec.get("doctype") or []))
        put_text(ws, r, 14, old.get("snap"))
        put_text(ws, r, 15, doctype_sources(rec))
        put_text(ws, r, 16, doctype_stratum(rec))
        put_text(ws, r, 17, x["round"])
        ws.cell(r, 18, doctype_formula(r, "M"))
        # 1차 스냅샷 채점은 1차 행에서만 뜻이 있다 — 새 행은 스냅샷 라벨이 없다.
        ws.cell(r, 19, f'=IF($Q{r}="1차",{doctype_formula(r, "N")[1:]},"")')
    for col in "LMNOPQRS":
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:K{last}"
    dv1 = DataValidation(type="list", formula1="분류체계!$A$2:$A$28", allow_blank=True)
    dv1.add(f"H2:I{last}")
    dv2 = DataValidation(type="list", formula1='"정답있음,분류대상아님,체계에없음,판단불가"', allow_blank=True)
    dv2.add(f"J2:J{last}")
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)

    # 분류체계는 1차 시트의 목록을 그대로 쓴다(드롭다운 원본 — 27개 노드).
    old_tax = openpyxl.load_workbook(OLD_XLSX)["분류체계"]
    for row in old_tax.iter_rows(values_only=True):
        tax.append(list(row))
    for c in range(1, 5):
        tax.cell(1, c).fill, tax.cell(1, c).font = HDR_FILL, HDR_FONT
    tax.column_dimensions["A"].width = 14
    tax.column_dimensions["B"].width = 36

    R = f"평가셋!$R$2:$R${last}"
    P = f"평가셋!$P$2:$P${last}"
    Q = f"평가셋!$Q$2:$Q${last}"
    J = f"평가셋!$J$2:$J${last}"
    a = [["업무분류 평가셋 — 집계"],
         ["이 시트는 전부 수식이다. 평가셋 시트에 정답을 채우면 숫자가 따라 바뀐다. 비어 보이면 Ctrl+Alt+F9."],
         [],
         ["■ 라벨링 진행"],
         ["표본 총건수", f"=COUNT(평가셋!A2:A{last})"],
         ["판정 입력 완료", f"=COUNTA(평가셋!J2:J{last})"],
         ["미입력", "=B5-B6"],
         ["진행률", "=IF(B5=0,0,B6/B5)"],
         [],
         ["■ 판정 분포", "건수", "비율"]]
    for i, v in enumerate(("정답있음", "분류대상아님", "체계에없음", "판단불가")):
        rr = 11 + i
        a.append([v, f'=COUNTIF({J},A{rr})', f"=IF($B$6=0,0,B{rr}/$B$6)"])
    a += [[], ["■ 성능 — 표본 기준", "적중", "오라벨", "헛라벨", "놓침", "정상무라벨", "정밀도", "재현율"]]
    specs = [("전체 표본 · 현행 엔진", R, None), (f"1차 {n_old}건 · 현행 엔진", R, "1차"),
             (f"1차 {n_old}건 · 0826 스냅샷", f"평가셋!$S$2:$S${last}", None)]
    for i, (label, rng_, rnd) in enumerate(specs):
        rr = 17 + i
        cells = []
        for k in ("적중", "오라벨", "헛라벨", "놓침", "정상무라벨"):
            cells.append(f'=COUNTIFS({Q},"1차",{rng_},"{k}")' if rnd else f'=COUNTIF({rng_},"{k}")')
        a.append([label] + cells + [f"=IF(B{rr}+C{rr}+D{rr}=0,0,B{rr}/(B{rr}+C{rr}+D{rr}))",
                                    f"=IF(B{rr}+C{rr}+E{rr}=0,0,B{rr}/(B{rr}+C{rr}+E{rr}))"])
    a += [["정밀도 = 적중/(적중+오라벨+헛라벨) · 재현율 = 적중/(적중+오라벨+놓침)"], [],
          ["■ 층별 가중치 — 층마다 뽑은 비율이 달라 표본 평균 ≠ 모집단 평균", "모집단 N", "표본 n", "가중치 N/n"]]
    w0 = len(a) + 1
    for i, s in enumerate(DOC_STRATA):
        rr = w0 + i
        a.append([s, pop[s], f'=COUNTIF({P},A{rr})', f"=IF(C{rr}=0,0,B{rr}/C{rr})"])
    w1 = w0 + len(DOC_STRATA) - 1
    a.append(["합계", f"=SUM(B{w0}:B{w1})", f"=SUM(C{w0}:C{w1})"])
    a += [[], [f"■ 모집단 {sum(pop.values())}건 추정 — 보고에 쓸 숫자는 이쪽이다",
               "가중 적중", "가중 오라벨", "가중 헛라벨", "가중 놓침", "정밀도", "재현율"]]
    rr = len(a) + 1
    sp = [f'=SUMPRODUCT($D${w0}:$D${w1},COUNTIFS({P},$A${w0}:$A${w1},{R},"{k}"))'
          for k in ("적중", "오라벨", "헛라벨", "놓침")]
    a.append(["현행 엔진"] + sp + [f"=IF(B{rr}+C{rr}+D{rr}=0,0,B{rr}/(B{rr}+C{rr}+D{rr}))",
                                  f"=IF(B{rr}+C{rr}+E{rr}=0,0,B{rr}/(B{rr}+C{rr}+E{rr}))"])
    write_rows(agg, a)

    g = [["업무분류 평가셋 — 작성지침 (2차 확장, 2026-09-14)"],
         [],
         ["목적", "엔진이 붙인 업무분류 라벨이 맞는지(정밀도)와, 붙어야 할 문서에 붙었는지(재현율)를 사람이 판정해 잰다."],
         ["대상", f"D:\\분류함 중 본문을 읽을 수 있는 문서 {sum(pop.values())}건(내용이 같은 복사본은 1건으로)에서 층화 표집한 {len(rows)}건."],
         ["구성", f"1차(2026-08-26) 사람 라벨 {n_old}건을 그대로 옮기고, {len(rows) - n_old}건을 새로 뽑았다(ref_회차 열)."],
         ["엔진", f"Rust 판 --rule-only · 규칙 {meta.get('doctype_rule_version')} · 분류체계 {meta.get('taxonomy_version')}"],
         [],
         ["■ 채울 칸 (노란색) — 2차 행(no 가 뒤쪽)만 채우면 된다. 1차 행은 이미 채워져 있다"],
         ["정답1_dc_id", "이 문서의 올바른 업무분류. '분류체계' 시트 목록에서 고른다(드롭다운)."],
         ["정답2_dc_id", "라벨이 2개여야 한다고 판단될 때만. 아니면 비워 둔다."],
         ["판정", "정답있음 / 분류대상아님 / 체계에없음 / 판단불가 중 하나(드롭다운)."],
         ["비고", "판단 근거나 애매했던 이유를 짧게. 선택."],
         [],
         ["■ 판정 값의 뜻"],
         ["정답있음", "27개 분류체계 중 맞는 항목이 있다 → 정답1을 반드시 채운다."],
         ["분류대상아님", "업무문서가 아니다(신문기사·논문·판례·공개 데이터셋 등). 라벨이 없는 게 맞다. 정답칸은 비운다."],
         ["체계에없음", "업무문서는 맞는데 27개 항목 어디에도 해당이 없다(예: 구입의뢰서, 점검확인서). 정답칸은 비운다."],
         ["판단불가", "발췌만으로 알 수 없고 원본을 열어봐도 애매하다."],
         [],
         ["■ 중요 — 이렇게 하지 마세요"],
         ["엔진 답 보지 않기", "엔진이 뭐라고 했는지는 숨김열(L~S)에 있다. 라벨링이 끝나기 전에는 열지 마세요."],
         ["", "보고 나서 붙이면 규칙을 검증하는 게 아니라 규칙을 베끼는 것이 되어 결과가 무의미해진다."],
         ["채우기 끌기 조심", "1차에서 판정이 '분류대상아님'인데 정답1이 채워진 행 23건이 끌기 흔적으로 나왔다. 판정과 정답칸이 맞는지 확인."],
         ["공백 조심", "dc_id 뒤에 공백이 붙으면 불일치로 세어진다(1차에서 14건). 드롭다운으로 고르면 생기지 않는다."],
         [],
         ["■ 표집 방법"],
         ["층", "엔진이 붙인 라벨 수로 D1 라벨1개 · D2 라벨2개이상 · D3 미라벨. 오탐은 D2 에, 놓침은 D3 에 몰린다."],
         ["배분", "1차 행을 층에 먼저 넣고, 층별 목표(√모집단 비례)에서 모자란 만큼 새로 뽑았다. 층 안에서는 포맷 비례."],
         ["주의", "층마다 뽑은 비율이 달라, 단순 평균이 아니라 '집계' 시트의 모집단 추정치를 봐야 한다."],
         ["", "1차 행은 1차 모집단(354건)에서 뽑혔던 문서라 2차 모집단에 대해 완전한 무작위 표본은 아니다 — 집계에 1차만 따로 본 줄을 두었다."],
         ["재현", f"시드 {SEED} 고정 — tests/evalset/build_evalsets.py 를 같은 입력으로 다시 돌리면 같은 표본이 나온다."],
         [],
         ["■ 열 구성"],
         ["A~G", "문서 정보(폴더·파일명·포맷·본문 글자 수·첫 줄·앞부분 300자 발췌)"],
         ["H·I·J·K", "사람이 채우는 칸 — 정답1 · 정답2 · 판정 · 비고"],
         ["L~Q (숨김)", "상대경로 · 현행 엔진 라벨 · 1차 시점 엔진 라벨(1차 행만) · 신호원 · 층 · 회차"],
         ["R·S (숨김)", "채점 결과(수식) — 적중 · 오라벨 · 헛라벨 · 놓침 · 정상무라벨"],
         [],
         ["■ 1차에서 옮기지 못한 행", f"{len(dropped)}건"]]
    for name, why in dropped:
        g.append(["", f"{name} — {why}"])
    write_rows(guide, g)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 120
    wb.save(path)


#------------------------------------------------------------------
# 보안등급 평가셋 xlsx 쓰기
#=> 시트 셋: 작성지침 · 평가셋 · 집계(전부 수식). 엔진 등급·근거·채점은 숨김열.
#   집계는 C 재현율·C 정밀도·누출 위험(정답 C 인데 엔진 O·미분류)·혼동행렬과
#   층별 가중치로 되돌린 모집단 추정을 낸다.
#
# -in: rows = pick_security 의 표본 행
# -in: pop  = {층: 모집단 수}
# -in: meta = {rule_version, ...}
# -in: path = 저장 경로
#
# -out: 없음(파일을 쓴다)
# -out: error = 저장 실패 시 OSError 전파
#------------------------------------------------------------------
def write_security_book(rows, pop, meta, path):
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "작성지침"
    ws = wb.create_sheet("평가셋")
    agg = wb.create_sheet("집계")
    last = len(rows) + 1

    headers = ["no", "폴더", "파일명", "포맷", "규모", "첫 줄", "앞부분 발췌", "정답등급", "근거유형",
               "비고", "ref_상대경로", "ref_엔진등급", "ref_method", "ref_근거신호", "ref_층", "결과"]
    write_header(ws, headers, {"A": 5, "B": 20, "C": 42, "D": 7, "E": 10, "F": 30, "G": 80,
                               "H": 11, "I": 16, "J": 26, "K": 30, "L": 8, "M": 12, "N": 40,
                               "O": 14, "P": 10})
    for i, x in enumerate(rows, start=1):
        r = i + 1
        d, rec = x["doc"], x["doc"]["rec"]
        first, exc = excerpt(d["text"], 600)
        ws.cell(r, 1, i)
        put_text(ws, r, 2, rel_folder(d["path"]))
        put_text(ws, r, 3, os.path.basename(d["path"]))
        put_text(ws, r, 4, d["ext"])
        put_text(ws, r, 5, f"{d['chars']:,}자")
        put_text(ws, r, 6, first)
        put_text(ws, r, 7, exc).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (8, 9, 10):
            ws.cell(r, col).fill = IN_FILL
        put_text(ws, r, 11, os.path.relpath(d["path"], SRC_DIR))
        put_text(ws, r, 12, rec.get("grade"))
        put_text(ws, r, 13, ((rec.get("why") or {}).get("security") or {}).get("method"))
        put_text(ws, r, 14, security_signals(rec))
        put_text(ws, r, 15, security_stratum(rec))
        # 정답 C 를 엔진이 C 로 잡았나(적중)·못 잡았나(놓침), 정답이 C 가 아닌데 C 로 올렸나(오탐).
        ws.cell(r, 16, f'=IF($H{r}="","",IF($H{r}="판단불가","판단불가",IF($H{r}="C",'
                       f'IF($L{r}="C","C적중","C놓침"),IF($L{r}="C","C오탐","비C정상"))))')
    for col in "KLMNOP":
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:J{last}"
    dv1 = DataValidation(type="list", formula1='"C,S,O,판단불가"', allow_blank=True)
    dv1.add(f"H2:H{last}")
    dv2 = DataValidation(type="list", formula1=f'"{SEC_BASIS}"', allow_blank=True)
    dv2.add(f"I2:I{last}")
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)

    H = f"평가셋!$H$2:$H${last}"
    L = f"평가셋!$L$2:$L${last}"
    O = f"평가셋!$O$2:$O${last}"
    Pr = f"평가셋!$P$2:$P${last}"
    a = [["보안등급 평가셋 — 집계 (C 재현율 중심)"],
         ["이 시트는 전부 수식이다. 평가셋 시트에 정답을 채우면 숫자가 따라 바뀐다. 비어 보이면 Ctrl+Alt+F9."],
         [],
         ["■ 라벨링 진행"],
         ["표본 총건수", f"=COUNT(평가셋!A2:A{last})"],
         ["정답 입력 완료", f"=COUNTA(평가셋!H2:H{last})"],
         ["미입력", "=B5-B6"],
         ["진행률", "=IF(B5=0,0,B6/B5)"],
         [],
         ["■ 정답 분포", "건수", "비율"]]
    for i, v in enumerate(("C", "S", "O", "판단불가")):
        rr = 11 + i
        a.append([v, f"=COUNTIF({H},A{rr})", f"=IF($B$6=0,0,B{rr}/$B$6)"])
    a += [[], ["■ C 성능 — 표본 기준", "C적중", "C놓침", "C오탐", "비C정상", "C 재현율", "C 정밀도"],
          ["현행 엔진", f'=COUNTIF({Pr},"C적중")', f'=COUNTIF({Pr},"C놓침")', f'=COUNTIF({Pr},"C오탐")',
           f'=COUNTIF({Pr},"비C정상")', "=IF(B17+C17=0,0,B17/(B17+C17))", "=IF(B17+D17=0,0,B17/(B17+D17))"],
          ["C 재현율 = C적중/(C적중+C놓침) — 진짜 C 중 엔진이 C 로 잡은 비율 · C 정밀도 = C적중/(C적중+C오탐)"],
          [],
          ["■ 누출 위험 — 정답 C 인데 엔진이 낮춘 문서(가장 치명적)", "건수"],
          ["정답 C · 엔진 O", f'=COUNTIFS({H},"C",{L},"O")'],
          ["정답 C · 엔진 미분류", f'=COUNTIFS({H},"C",{L},"")'],
          ["정답 C · 엔진 S", f'=COUNTIFS({H},"C",{L},"S")'],
          [],
          ["■ 혼동행렬 — 행: 정답 · 열: 엔진", "엔진 C", "엔진 S", "엔진 O", "엔진 미분류"]]
    for i, g_ in enumerate(("C", "S", "O")):
        a.append([f"정답 {g_}"] + [f'=COUNTIFS({H},"{g_}",{L},"{e}")' for e in ("C", "S", "O", "")])
    a += [[], ["■ 층별 가중치 — 층마다 뽑은 비율이 달라 표본 평균 ≠ 모집단 평균", "모집단 N", "표본 n", "가중치 N/n"]]
    w0 = len(a) + 1
    for i, s in enumerate(SEC_STRATA):
        rr = w0 + i
        a.append([s, pop[s], f"=COUNTIF({O},A{rr})", f"=IF(C{rr}=0,0,B{rr}/C{rr})"])
    w1 = w0 + len(SEC_STRATA) - 1
    a.append(["합계", f"=SUM(B{w0}:B{w1})", f"=SUM(C{w0}:C{w1})"])
    a += [[], [f"■ 모집단 {sum(pop.values())}건 추정 — 보고에 쓸 숫자는 이쪽이다",
               "가중 C적중", "가중 C놓침", "가중 C오탐", "C 재현율", "C 정밀도"]]
    rr = len(a) + 1
    sp = [f'=SUMPRODUCT($D${w0}:$D${w1},COUNTIFS({O},$A${w0}:$A${w1},{Pr},"{k}"))'
          for k in ("C적중", "C놓침", "C오탐")]
    a.append(["현행 엔진"] + sp + [f"=IF(B{rr}+C{rr}=0,0,B{rr}/(B{rr}+C{rr}))",
                                  f"=IF(B{rr}+D{rr}=0,0,B{rr}/(B{rr}+D{rr}))"])
    write_rows(agg, a)

    g = [["보안등급 평가셋 — 작성지침 (C 재현율, 2026-09-14)"],
         [],
         ["목적", "진짜 기밀(C) 문서 중 엔진이 C 로 잡은 비율(C 재현율)을 잰다. C 를 O(공개)로 내보내는 것이 가장 치명적인 실수라서다."],
         ["대상", f"D:\\분류함 중 본문을 읽을 수 있는 문서 {sum(pop.values())}건(내용이 같은 복사본은 1건으로)에서 층화 표집한 {len(rows)}건."],
         ["엔진", f"Rust 판 --rule-only · 보안 규칙 {meta.get('rule_version')}"],
         ["주의", "발췌에 실제 문서 내용(개인정보 포함 가능)이 들어 있다. 이 파일을 외부로 보내거나 공유 폴더에 두지 마세요."],
         [],
         ["■ 채울 칸 (노란색)"],
         ["정답등급", "C / S / O / 판단불가 중 하나(드롭다운). 발췌만으로 애매하면 원본(숨김열 K 가 아니라 폴더·파일명으로 찾아)을 열어 본다."],
         ["근거유형", "그 등급으로 본 가장 큰 이유 하나(드롭다운). C·S 일 때는 꼭 채운다 — 놓친 원인을 가르는 데 쓴다."],
         ["비고", "애매했던 점, 원본을 열어 봤는지 등. 선택."],
         [],
         ["■ 등급 판정 기준 — 초안. 사내 보안정책이 우선하며, 담당자가 확정한 뒤 라벨링을 시작한다"],
         ["C (기밀)", "유출되면 회사나 개인에게 심각한 피해가 나는 문서."],
         ["", "· 개인을 특정할 수 있는 정보(주민등록번호·계좌·카드번호, 이름+연락처+주소 결합 등)"],
         ["", "· 법상 민감정보(건강·범죄경력·사상/신념·노조 등 — 개인정보보호법 제23조)"],
         ["", "· 영업·기술 기밀(단가·견적·원가, 미공개 설계·소스, 고객사 내부 구성), 계약서, 인사·급여"],
         ["", "· '대외비·극비·CONFIDENTIAL' 표시가 있는 문서"],
         ["S (민감)", "외부 공개는 곤란하지만 유출 피해가 제한적인 사내 업무자료(내부 매뉴얼·제안서 초안·회의자료 등)."],
         ["O (공개)", "이미 공개됐거나 공개해도 되는 자료(신문기사·공개 법령/판례·공개 가이드·논문·공개 데이터셋·배포용 홍보물)."],
         ["판단불가", "원본을 열어 봐도 등급을 정할 수 없을 때만. 남발하면 재현율을 잴 수 없다."],
         ["경계에서", "C 와 S 사이에서 망설여지면 '유출되면 누가 얼마나 다치나'로 가른다. 망설인 이유는 비고에 적는다."],
         [],
         ["■ 중요 — 이렇게 하지 마세요"],
         ["엔진 답 보지 않기", "엔진 등급·근거는 숨김열(K~P)에 있다. 라벨링이 끝나기 전에는 열지 마세요. 보고 나서 달면 베끼기가 된다."],
         ["파일명만 보고 정하기", "파일명이 무의미한 문서가 많다(1차 업무분류 검증에서 확인). 발췌를 읽고 정한다."],
         [],
         ["■ 표집 방법"],
         ["층", "엔진 등급으로 G1 엔진C · G2 엔진S · G3 엔진O · G4 엔진미분류. 놓친 C 는 G2~G4 에 숨어 있어 일부러 넉넉히 넣었다."],
         ["배분", "층마다 최소 60건을 먼저 주고 나머지를 √모집단 비례로 나눴다. 층 안에서는 포맷 비례."],
         ["주의", "층마다 뽑은 비율이 달라, 단순 평균이 아니라 '집계' 시트의 모집단 추정치를 봐야 한다."],
         ["범위", "스캔 PDF(처리제외)·본문을 못 읽은 문서는 넣지 않았다 — 발췌로 판정할 수 없고, 1차 범위 밖이다."],
         ["재현", f"시드 {SEED} 고정 — tests/evalset/build_evalsets.py 를 같은 입력으로 다시 돌리면 같은 표본이 나온다."],
         [],
         ["■ 열 구성"],
         ["A~G", "문서 정보(폴더·파일명·포맷·본문 글자 수·첫 줄·앞부분 600자 발췌)"],
         ["H·I·J", "사람이 채우는 칸 — 정답등급 · 근거유형 · 비고"],
         ["K~O (숨김)", "상대경로 · 엔진 등급 · 판정 방식 · 근거 신호 · 층"],
         ["P (숨김)", "채점 결과(수식) — C적중 · C놓침 · C오탐 · 비C정상 · 판단불가"]]
    write_rows(guide, g)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 120
    wb.save(path)


#------------------------------------------------------------------
# 진입점
#=> 문서 목록 → 엔진 실행(캐시) → 결과 붙이기 → 두 평가셋 표집·저장 → 요약 출력.
#
# -in: 없음(명령줄: --work, --doctype-total, --security-total)
#
# -out: 없음(xlsx 두 개와 작업폴더의 manifest.json 을 쓴다)
# -out: error = 엔진 exe 가 없으면 SystemExit, 그 밖의 오류는 전파
#------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="업무분류·보안등급 평가셋 생성")
    ap.add_argument("--work", default=os.path.join(tempfile.gettempdir(), "csoclassify_evalset"))
    ap.add_argument("--doctype-total", type=int, default=500)
    ap.add_argument("--security-total", type=int, default=500)
    args = ap.parse_args()
    os.makedirs(args.work, exist_ok=True)

    pop_file = os.path.join(args.work, "population_files.json")
    if os.path.isfile(pop_file):
        docs = json.load(open(pop_file, encoding="utf-8"))
    else:
        docs = enumerate_docs(SRC_DIR)
        json.dump(docs, open(pop_file, "w", encoding="utf-8"), ensure_ascii=False)
    run_engine(args.work, docs)
    attach_results(args.work, docs)

    meta = {}
    for line in open(os.path.join(args.work, "run.jsonl"), encoding="utf-8"):
        if line.startswith('{"run"'):
            meta = json.loads(line)["run"]
            break

    d_rows, d_pop, dropped = pick_doctype(docs, args.doctype_total)
    s_rows, s_pop = pick_security(docs, args.security_total)
    write_doctype_book(d_rows, d_pop, dropped, meta, OUT_DOCTYPE)
    write_security_book(s_rows, s_pop, meta, OUT_SECURITY)

    manifest = {
        "seed": SEED, "meta": meta,
        "documents": len(docs), "usable": sum(1 for d in docs if d["usable"]),
        "doctype": {"population": d_pop, "sample": collections.Counter(doctype_stratum(x["doc"]["rec"]) for x in d_rows),
                    "rounds": collections.Counter(x["round"] for x in d_rows),
                    "formats": collections.Counter(x["doc"]["ext"] for x in d_rows),
                    "dropped": dropped, "hashes": [x["doc"]["hash"] for x in d_rows]},
        "security": {"population": s_pop, "sample": collections.Counter(security_stratum(x["doc"]["rec"]) for x in s_rows),
                     "formats": collections.Counter(x["doc"]["ext"] for x in s_rows),
                     "hashes": [x["doc"]["hash"] for x in s_rows]},
    }
    json.dump(manifest, open(os.path.join(args.work, "manifest.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(json.dumps({k: v for k, v in manifest.items() if k not in ("doctype", "security")}, ensure_ascii=False))
    for k in ("doctype", "security"):
        print(k, json.dumps({x: y for x, y in manifest[k].items() if x != "hashes"}, ensure_ascii=False))
    print("저장:", OUT_DOCTYPE)
    print("저장:", OUT_SECURITY)


if __name__ == "__main__":
    main()
